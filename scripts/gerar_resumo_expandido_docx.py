#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
===============================================================================
Gerador do Resumo Expandido — Seminário PROPPG/UTFPR (SBPROPPG 2026)
Programa de Pós-Graduação em Tecnologias Computacionais para o Agronegócio (PPGTCA)
Pesquisa: Geolocalização e Predição de Erosão Laminar — SAREL v2.0
===============================================================================

LIMITES E DIRETRIZES TÉCNICAS:
- Limite de Extensão: O conjunto das seções principais (Título, Autores, Introdução,
  Objetivos, Metodologia, Resultados e Considerações Finais) deve conter NO MÁXIMO 500 palavras.
  A seção de Referências não é contabilizada.
- Normas: Formatação científica em estrita observância à APA (7ª Edição).
- Modelo Base: docs/Seminario prov/Modelo_resumo_SBPROPPG_26.docx (preserva banner/header).
"""

import os
import sys
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def criar_resumo_expandido():
    caminho_modelo = os.path.join("docs", "Seminario prov", "Modelo_resumo_SBPROPPG_26.docx")
    caminho_saida = os.path.join("docs", "Seminario prov", "Resumo_Expandido_SBPROPPG_26_Luis_Alfredo.docx")

    if not os.path.exists(caminho_modelo):
        raise FileNotFoundError(f"Arquivo modelo não encontrado: {caminho_modelo}")

    doc = docx.Document(caminho_modelo)

    # Limpa os parágrafos do corpo mantendo os cabeçalhos, imagens e seções
    p_elements = list(doc.element.body)
    for elem in p_elements:
        if elem.tag.endswith('p') or elem.tag.endswith('tbl'):
            doc.element.body.remove(elem)

    # =========================================================================
    # Definição dos Textos Científicos com Controle Rigoroso de Palavras
    # =========================================================================
    titulo_texto = (
        "ARCABOUÇO COMPUTACIONAL AUDITÁVEL PARA AMOSTRAGEM ESTRATIFICADA E PREDIÇÃO "
        "DE EROSÃO LAMINAR EM BACIA AGRÍCOLA VIA SENSORIAMENTO MULTIESPECTRAL E APRENDIZADO DE MÁQUINA"
    )

    autores_texto = "Luis Alfredo da Silva¹; Coautor²; Nome do Orientador³"
    filiacao_texto = (
        "¹Mestrando do Programa de Pós-Graduação em Tecnologias Computacionais para o Agronegócio (PPGTCA), "
        "Universidade Tecnológica Federal do Paraná (UTFPR). ²Coautor. ³Docente Orientador, PPGTCA/UTFPR."
    )

    intro_corpo = (
        "A erosão laminar degrada a fertilidade e reduz a produtividade na Bacia do Paraná 3, sob Sistema Plantio "
        "Direto (Dieckow et al., 2009). Modelos empíricos como a RUSLE apoiam-se em generalizações cartográficas que "
        "não capturam a heterogeneidade das feições em escala métrica. A integração entre constelações orbitais "
        "multiespectrais, veículos aéreos não tripulados (VANT) de alta resolução e aprendizado de máquina "
        "supervisionado abre caminho para a detecção e a predição espacialmente explícitas de focos ativos, "
        "reduzindo custos de campo e viabilizando o manejo conservacionista de precisão."
    )

    objetivo_corpo = (
        "Desenvolver e validar um arcabouço computacional auditável (SAREL) para amostragem estratificada, calibração "
        "radiométrica multiescala e predição de risco de erosão laminar no Oeste do Paraná, integrando séries "
        "temporais do Sentinel-2 MSI, ortomosaicos centimétricos de VANT e o algoritmo XGBoost sob validação espacial em blocos."
    )

    metodologia_corpo = (
        "O desenho amostral organiza-se em 18 estratos biofísicos ortogonais de declividade (3% a 20%), frequência "
        "multitemporal de solo exposto (índice BSI) e erodibilidade pedológica oficial (Coelho et al., 2024). "
        "Aplica-se thinning geodésico Haversine (1 km a 5 km) sobre imóveis do SICAR, mitigando a dependência "
        "espacial (Tobler, 1970). Atributos topográficos são reduzidos do Copernicus DEM GLO-30 em projeção métrica "
        "SIRGAS 2000 / UTM 22S (EPSG:31982), e o Fator C da RUSLE deriva de séries NDVI (Durigon et al., 2014). "
        "A validação independente (held-out) prevê quatro sítios contínuos de 10 a 50 ha, a imagear por VANT "
        "multiespectral (cinco bandas, RTK/PPK, GSD 3 a 7,5 cm) e confrontar radiometricamente com o Sentinel-2. "
        "O treinamento adotará XGBoost (Chen & Guestrin, 2016) sob validação cruzada Leave-One-Catchment-Out "
        "(Roberts et al., 2017), com explicabilidade via SHAP (Lundberg & Lee, 2017)."
    )

    resultados_corpo = (
        "Esta etapa reporta a validação da infraestrutura computacional, não do desempenho preditivo. O sistema "
        "sustenta 270 testes automatizados e um varredor estático que barra constantes arbitrárias e dados "
        "sintéticos: toda variável carrega proveniência e data de aquisição, e ausência de dado nunca vira valor. "
        "Implementou-se a arquitetura dual, que segrega a detecção de focos ativos (Modelo D, t0) da predição "
        "prospectiva (Modelo P), com guarda temporal de 24 meses que previne vazamento de dados (Kaufman et al., "
        "2012) e espelha a rotação de culturas. Executou-se um benchmark de integridade do fluxo XGBoost + SHAP sob "
        "partição espacial, com a classe de controle sintetizada apenas para teste do pipeline e assim sinalizada "
        "no relatório: suas métricas aferem o código, não a separabilidade do fenômeno."
    )

    consideracoes_corpo = (
        "O arcabouço mostrou-se reprodutível e pericialmente auditável, com rastreabilidade ponto a ponto — "
        "contribuição direta para perícias agronômicas e governança territorial de bacias. A limitação principal é "
        "que a campanha de rotulagem e os sobrevoos ainda não foram executados; o desempenho preditivo será aferido "
        "após os rótulos periciais."
    )

    # Auditoria de Extensão (Limite de 500 palavras)
    secoes_principais = [
        titulo_texto, autores_texto, filiacao_texto,
        f"Introdução: {intro_corpo}",
        f"Objetivo: {objetivo_corpo}",
        f"Metodologia: {metodologia_corpo}",
        f"Resultados e Discussão: {resultados_corpo}",
        f"Considerações Finais: {consideracoes_corpo}"
    ]
    total_palavras = sum(len(s.split()) for s in secoes_principais)
    print(f"[AUDITORIA DE EXTENSÃO] Total de palavras das seções principais: {total_palavras} (Limite máximo: 500)")
    if total_palavras > 500:
        raise ValueError(f"Extensão excedida: {total_palavras} palavras ultrapassa o limite de 500.")

    # =========================================================================
    # Inserção e Formatação no Documento Word (.docx)
    # =========================================================================

    # 1. Título
    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tit.paragraph_format.space_before = Pt(0)
    p_tit.paragraph_format.space_after = Pt(6)
    r_tit = p_tit.add_run(titulo_texto)
    r_tit.bold = True
    r_tit.font.name = "Calibri"
    r_tit.font.size = Pt(12)

    # 2. Autores
    p_aut = doc.add_paragraph()
    p_aut.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aut.paragraph_format.space_before = Pt(0)
    p_aut.paragraph_format.space_after = Pt(2)
    r_aut = p_aut.add_run(autores_texto)
    r_aut.font.name = "Calibri"
    r_aut.font.size = Pt(10.5)

    # Filiação Institucional
    p_fil = doc.add_paragraph()
    p_fil.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_fil.paragraph_format.space_before = Pt(0)
    p_fil.paragraph_format.space_after = Pt(8)
    r_fil = p_fil.add_run(filiacao_texto)
    r_fil.font.name = "Calibri"
    r_fil.font.size = Pt(9)
    r_fil.font.italic = True

    # Função auxiliar para seções com título em negrito seguido de dois-pontos
    def adicionar_secao(titulo_secao, texto_secao):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15

        r_rot = p.add_run(f"{titulo_secao}: ")
        r_rot.bold = True
        r_rot.font.name = "Calibri"
        r_rot.font.size = Pt(11)

        r_txt = p.add_run(texto_secao)
        r_txt.font.name = "Calibri"
        r_txt.font.size = Pt(11)
        return p

    # 3. Introdução
    adicionar_secao("Introdução", intro_corpo)

    # 4. Objetivo
    adicionar_secao("Objetivo", objetivo_corpo)

    # 5. Metodologia
    adicionar_secao("Metodologia", metodologia_corpo)

    # 6. Resultados e Discussão
    adicionar_secao("Resultados e Discussão", resultados_corpo)

    # 7. Considerações Finais
    adicionar_secao("Considerações Finais", consideracoes_corpo)

    # 8. Palavras-chave
    p_kw = doc.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kw.paragraph_format.space_before = Pt(6)
    p_kw.paragraph_format.space_after = Pt(6)
    p_kw.paragraph_format.line_spacing = 1.15
    r_kw_label = p_kw.add_run("Palavras-chave: ")
    r_kw_label.bold = True
    r_kw_label.font.name = "Calibri"
    r_kw_label.font.size = Pt(11)
    r_kw_txt = p_kw.add_run("erosão laminar, sensoriamento remoto, aprendizado de máquina, vant multiespectral, bacia hidrográfica.")
    r_kw_txt.font.name = "Calibri"
    r_kw_txt.font.size = Pt(11)

    # 9. Agradecimentos (Nota do Autor)
    p_agr_lbl = doc.add_paragraph()
    p_agr_lbl.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_agr_lbl.paragraph_format.space_before = Pt(6)
    p_agr_lbl.paragraph_format.space_after = Pt(2)
    r_agr_lbl = p_agr_lbl.add_run("Agradecimentos (Nota do Autor):")
    r_agr_lbl.bold = True
    r_agr_lbl.font.name = "Calibri"
    r_agr_lbl.font.size = Pt(11)

    p_agr_txt = doc.add_paragraph()
    p_agr_txt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_agr_txt.paragraph_format.space_before = Pt(0)
    p_agr_txt.paragraph_format.space_after = Pt(8)
    p_agr_txt.paragraph_format.line_spacing = 1.15
    r_agr_val = p_agr_txt.add_run(
        "O presente trabalho foi realizado com o apoio do Programa de Pós-Graduação em Tecnologias Computacionais "
        "para o Agronegócio (PPGTCA) da Universidade Tecnológica Federal do Paraná (UTFPR) e do Itaipu Parquetec."
    )
    r_agr_val.font.name = "Calibri"
    r_agr_val.font.size = Pt(10)

    # 10. Referências (Normas APA 7ª Edição)
    p_ref_lbl = doc.add_paragraph()
    p_ref_lbl.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_ref_lbl.paragraph_format.space_before = Pt(8)
    p_ref_lbl.paragraph_format.space_after = Pt(4)
    r_ref_lbl = p_ref_lbl.add_run("Referências")
    r_ref_lbl.bold = True
    r_ref_lbl.font.name = "Calibri"
    r_ref_lbl.font.size = Pt(11)

    referencias_apa = [
        {
            "autores": "Chen, T., & Guestrin, C. (2016). ",
            "titulo": "XGBoost: A scalable tree boosting system. ",
            "fonte_italico": "Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining, ",
            "complemento": "785–794. https://doi.org/10.1145/2939672.2939785"
        },
        {
            "autores": ("Coelho, M. R., Lumbreras, J. F., Amaral, A. J. do, Vasques, G. M., Mansilla Baca, J. F., "
                        "Dart, R. de O., & Pedreira, J. P. das N. C. (2024). "),
            "titulo": "",
            "fonte_italico": "Erodibilidade dos solos do Brasil ",
            "complemento": "(Documentos 246). Embrapa Solos. https://www.infoteca.cnptia.embrapa.br/infoteca/handle/doc/1170044"
        },
        {
            "autores": ("Dieckow, J., Bayer, C., Conceição, P. C., Zanatta, J. A., Martin-Neto, L., Milori, D. M. B. P., "
                        "Salton, J. C., Macedo, M. M., Mielniczuk, J., & Hernani, L. C. (2009). "),
            "titulo": "Land use, tillage, texture and organic matter stock and composition in tropical and subtropical Brazilian soils. ",
            "fonte_italico": "European Journal of Soil Science, ",
            "complemento": "60(2), 240–249. https://doi.org/10.1111/j.1365-2389.2008.01101.x"
        },
        {
            "autores": "Durigon, V. L., Carvalho, D. F., Antunes, M. A. H., Oliveira, P. T. S., & Fernandes, M. M. (2014). ",
            "titulo": "NDVI time series for monitoring RUSLE cover management factor in a tropical watershed. ",
            "fonte_italico": "International Journal of Remote Sensing, ",
            "complemento": "35(2), 441–453. https://doi.org/10.1080/01431161.2013.871081"
        },
        {
            "autores": "Kaufman, S., Rosset, S., Perlich, C., & Stitelman, O. (2012). ",
            "titulo": "Leakage in data mining: Formulation, detection, and avoidance. ",
            "fonte_italico": "ACM Transactions on Knowledge Discovery from Data, ",
            "complemento": "6(4), 1–21. https://doi.org/10.1145/2382577.2382579"
        },
        {
            "autores": "Lundberg, S. M., & Lee, S.-I. (2017). ",
            "titulo": "A unified approach to interpreting model predictions. In ",
            "fonte_italico": "Advances in Neural Information Processing Systems ",
            "complemento": ("(Vol. 30). Curran Associates. "
                            "https://papers.nips.cc/paper_files/paper/2017/hash/8a20a8621978632d76c43dfd28b67767-Abstract.html")
        },
        {
            "autores": ("Roberts, D. R., Bahn, V., Ciuti, S., Boyce, M. S., Elith, J., Guillera-Arroita, G., "
                        "Hauenstein, S., Lahoz-Monfort, J. J., Schröder, B., Thuiller, W., Warton, D. I., "
                        "Wintle, B. A., Hartig, F., & Dormann, C. F. (2017). "),
            "titulo": "Cross-validation strategies for data with temporal, spatial, hierarchical, or phylogenetic structure. ",
            "fonte_italico": "Ecography, ",
            "complemento": "40(8), 913–929. https://doi.org/10.1111/ecog.02881"
        },
        {
            "autores": "Tobler, W. R. (1970). ",
            "titulo": "A computer movie simulating urban growth in the Detroit region. ",
            "fonte_italico": "Economic Geography, ",
            "complemento": "46(sup1), 234–240. https://doi.org/10.2307/143141"
        }
    ]

    for ref in referencias_apa:
        p_ref = doc.add_paragraph()
        p_ref.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_ref.paragraph_format.space_before = Pt(2)
        p_ref.paragraph_format.space_after = Pt(4)
        p_ref.paragraph_format.line_spacing = 1.15

        # Recuo deslocado (hanging indent) de 1,27 cm (0.5 polegadas) padrão APA
        p_ref.paragraph_format.left_indent = Inches(0.5)
        p_ref.paragraph_format.first_line_indent = Inches(-0.5)

        r1 = p_ref.add_run(ref["autores"])
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)

        if ref["titulo"]:
            r2 = p_ref.add_run(ref["titulo"])
            r2.font.name = "Calibri"
            r2.font.size = Pt(10)

        if ref["fonte_italico"]:
            r3 = p_ref.add_run(ref["fonte_italico"])
            r3.font.name = "Calibri"
            r3.font.size = Pt(10)
            r3.italic = True

        if ref["complemento"]:
            r4 = p_ref.add_run(ref["complemento"])
            r4.font.name = "Calibri"
            r4.font.size = Pt(10)

    # Salva o arquivo preenchido
    doc.save(caminho_saida)
    print(f"[SUCESSO] Resumo expandido gerado com sucesso em:\n{caminho_saida}")
    return caminho_saida


if __name__ == "__main__":
    criar_resumo_expandido()
