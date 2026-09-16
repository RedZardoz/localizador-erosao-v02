# -*- coding: utf-8 -*-
"""
=============================================================================
GERADOR DA APRESENTAÇÃO EXECUTIVA E ACADÊMICA SAREL v2.0 (PPTX)
Programa de Pós-Graduação em Tecnologias Computacionais para o Agronegócio
PPGTCA 2026 - Pesquisa de Mestrado: Validação de Predição de Erosão Laminar
Autor: Luis Alfredo
=============================================================================
"""

import os, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Paleta de Cores Institucional do SAREL
C_NAVY_DARK = RGBColor(15, 23, 42)        # Slate 900 (#0F172A)
C_NAVY_CARD = RGBColor(30, 41, 59)        # Slate 800 (#1E293B)
C_BORDER_DARK = RGBColor(51, 65, 85)      # Slate 700 (#334155)

C_BG_LIGHT = RGBColor(248, 250, 252)      # Slate 50 (#F8FAFC)
C_BORDER_LIGHT = RGBColor(226, 232, 240)  # Slate 200 (#E2E8F0)
C_ACCENT_BG = RGBColor(236, 253, 245)     # Mint / Emerald 50 (#ECFDF5)

C_EMERALD_DARK = RGBColor(6, 95, 70)      # Emerald 800 (#065F46)
C_EMERALD_LIGHT = RGBColor(16, 185, 129)  # Emerald 500 (#10B981)
C_CYAN = RGBColor(14, 165, 233)           # Cyan 500 (#0EA5E9)
C_AMBER = RGBColor(217, 119, 6)           # Amber 600 (#D97706)
C_ROSE = RGBColor(225, 29, 72)            # Rose 600 (#E11D48)
C_INDIGO = RGBColor(99, 102, 241)         # Indigo 500 (#6366F1)

C_TEXT_DARK = RGBColor(15, 23, 42)        # Slate 900
C_TEXT_MUTED = RGBColor(71, 85, 105)      # Slate 600 (#475569)
C_TEXT_LIGHT = RGBColor(203, 213, 225)    # Slate 300 (#CBD5E1)
C_WHITE = RGBColor(255, 255, 255)

def set_slide_background(slide, color=C_BG_LIGHT):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, title_text, category_text="MESTRADO PPGTCA 2026 • TECNOLOGIAS COMPUTACIONAIS PARA O AGRONEGÓCIO"):
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = C_NAVY_DARK
    top_bar.line.color.rgb = C_NAVY_DARK

    tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.7), Inches(0.3))
    tf_cat = tb_cat.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(9.5)
    p_cat.font.bold = True
    p_cat.font.color.rgb = C_EMERALD_LIGHT

    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.38), Inches(11.7), Inches(0.65))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = C_WHITE

    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(1.15), Inches(13.333), Inches(0.05))
    line.fill.solid()
    line.fill.fore_color.rgb = C_EMERALD_LIGHT
    line.line.color.rgb = C_EMERALD_LIGHT

    ft = slide.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(11.7), Inches(0.3))
    p_ft = ft.text_frame.paragraphs[0]
    p_ft.text = "SAREL v2.0 • Sistema de Amostragem e Rotulagem para Erosão Laminar • PPGTCA 2026"
    p_ft.font.size = Pt(8.5)
    p_ft.font.color.rgb = C_TEXT_MUTED

def add_card(slide, left, top, width, height, title, content_lines, accent_color=C_EMERALD_LIGHT, bg_color=C_WHITE, border_color=C_BORDER_LIGHT, title_size=13, body_size=10):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.2)

    padding = Inches(0.2)
    tb = slide.shapes.add_textbox(left + padding, top + padding, width - (padding * 2), height - (padding * 2))
    tf = tb.text_frame
    tf.word_wrap = True

    if title:
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(title_size)
        p_title.font.bold = True
        p_title.font.color.rgb = accent_color
        p_title.space_after = Pt(5)

    first = not bool(title)
    for item in content_lines:
        p = tf.add_paragraph() if not first else tf.paragraphs[0]
        first = False
        if isinstance(item, tuple):
            strong_txt, norm_txt = item
            r1 = p.add_run()
            r1.text = strong_txt + " "
            r1.font.bold = True
            r1.font.size = Pt(body_size)
            r1.font.color.rgb = C_TEXT_DARK if bg_color != C_NAVY_CARD else C_WHITE
            r2 = p.add_run()
            r2.text = norm_txt
            r2.font.size = Pt(body_size)
            r2.font.color.rgb = C_TEXT_MUTED if bg_color != C_NAVY_CARD else C_TEXT_LIGHT
        else:
            p.text = str(item)
            p.font.size = Pt(body_size)
            p.font.color.rgb = C_TEXT_MUTED if bg_color != C_NAVY_CARD else C_TEXT_LIGHT
        p.space_after = Pt(4)

def add_metric_badge(slide, left, top, width, height, label, value, unit="", color=C_EMERALD_LIGHT, bg_color=C_WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = color
    shape.line.width = Pt(1.5)

    tb = slide.shapes.add_textbox(left, top + Inches(0.1), width, height - Inches(0.2))
    tf = tb.text_frame
    tf.word_wrap = True

    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    p1.text = label.upper()
    p1.font.size = Pt(9)
    p1.font.bold = True
    p1.font.color.rgb = C_TEXT_MUTED

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = f"{value} {unit}".strip()
    p2.font.size = Pt(17)
    p2.font.bold = True
    p2.font.color.rgb = color

def create_deck(output_pptx_path="Apresentacao_SAREL_PPGTCA_2026.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # =========================================================================
    # SLIDE 1: CAPA
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, C_NAVY_DARK)

    card_right = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(1.2), Inches(4.0), Inches(5.3))
    card_right.fill.solid()
    card_right.fill.fore_color.rgb = C_NAVY_CARD
    card_right.line.color.rgb = C_EMERALD_DARK

    tb_cr = s1.shapes.add_textbox(Inches(8.8), Inches(1.4), Inches(3.4), Inches(4.9))
    tf_cr = tb_cr.text_frame
    tf_cr.word_wrap = True
    
    p = tf_cr.paragraphs[0]
    p.text = "PILARES METODOLÓGICOS DO SAREL"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_EMERALD_LIGHT
    
    bullets_cr = [
        ("Desenho Amostral 18 Estratos", "Estratificação Ŝ x Ê x K̂ com Thinning Geodésico ≥ 1,0 km."),
        ("Lei Fundamental & Rigor Anti-Mock", "Dado verdadeiro ou ausência declarada. Zero dado fabricado."),
        ("Fusão Cadastral & LGPD", "559.899 imóveis SICAR-PR + SIGEF + SNCR com máscara byte a byte."),
        ("Validação Espacial LOCO", "XGBoost Leave-One-Catchment-Out particionado pelas bacias do PR."),
    ]
    for title, desc in bullets_cr:
        p1 = tf_cr.add_paragraph()
        p1.text = f"• {title}"
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = C_WHITE
        p1.space_before = Pt(8)
        p2 = tf_cr.add_paragraph()
        p2.text = desc
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = C_TEXT_LIGHT

    tb_main = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(7.2), Inches(5.0))
    tf_main = tb_main.text_frame
    tf_main.word_wrap = True

    p0 = tf_main.paragraphs[0]
    p0.text = "PROGRAMA DE PÓS-GRADUAÇÃO EM TECNOLOGIAS COMPUTACIONAIS PARA O AGRONEGÓCIO"
    p0.font.size = Pt(10)
    p0.font.bold = True
    p0.font.color.rgb = C_EMERALD_LIGHT

    p1 = tf_main.add_paragraph()
    p1.text = "SAREL: Sistema de Amostragem e Rotulagem para Erosão Laminar"
    p1.font.size = Pt(27)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE
    p1.space_before = Pt(12)

    p2 = tf_main.add_paragraph()
    p2.text = "Instrumento Computacional de Amostragem Multivariada, Sensoriamento Remoto e Gestão de Rotulagem Cega para Predição por Aprendizado de Máquina no Paraná"
    p2.font.size = Pt(13)
    p2.font.color.rgb = C_TEXT_LIGHT
    p2.space_before = Pt(12)

    p3 = tf_main.add_paragraph()
    p3.text = "Mestrando: Luis Alfredo | Pesquisa de Mestrado PPGTCA 2026\nLinha de Pesquisa: Tecnologias Geoespaciais e Modelagem Computacional no Agronegócio"
    p3.font.size = Pt(11)
    p3.font.bold = True
    p3.font.color.rgb = C_EMERALD_LIGHT
    p3.space_before = Pt(22)

    # =========================================================================
    # SLIDE 2: O PROBLEMA & A MOTIVAÇÃO CIENTÍFICA
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Contexto da Pesquisa: A Erosão Laminar e a Escala do Paraná")

    cols_s2 = [
        ("A Erosão Silenciosa", [
            ("Processo Invisível:", "A erosão laminar remove gradualmente as camadas superficiais férteis do solo agrícola sem abrir sulcos ou voçorocas imediatas."),
            ("Impacto Agronômico:", "Toneladas de matéria orgânica, fertilizantes e argila são carreadas aos rios muito antes do dano ser percebido visualmente."),
            ("Dimensão Territorial:", "O Estado do Paraná possui 199.315 km² de área, exigindo ferramentas capazes de monitorar bacias inteiras em macroescala."),
        ], C_AMBER),
        ("O Gargalo Tradicional", [
            ("Vistorias Dispersas:", "Inspeções de campo convencionais cobrem áreas restritas, são lentas e enviesadas por proximidade a rodovias e estradas."),
            ("Falta de Rastreabilidade:", "Planilhas manuais não registram a proveniência dos dados nem a data exata da imagem satelital que embasou o laudo."),
            ("Inviabilidade de Escala:", "Percorrer centenas de municípios paranaenses demanda tempo e recursos financeiros proibitivos para a pesquisa."),
        ], C_CYAN),
        ("A Proposta SAREL", [
            ("Estratificação Robusta:", "Algoritmo computacional que varre o território paranaense elegendo amostras estatisticamente balanceadas."),
            ("Fusão Multivariada:", "Integração do Google Earth Engine (Sentinel-2), Copernicus DEM, cartas oficiais da Embrapa e malhas do CAR."),
            ("Rigor Científico Total:", "Separação estrita entre observação e modelagem, garantindo blindagem metodológica perante a banca do PPGTCA."),
        ], C_EMERALD_LIGHT),
    ]

    for idx, (head, bullets, color) in enumerate(cols_s2):
        bx = Inches(0.8 + idx * 4.0)
        add_card(s2, bx, Inches(1.55), Inches(3.7), Inches(5.15), head, bullets, accent_color=color, title_size=14, body_size=10.2)

    # =========================================================================
    # SLIDE 3: O DILEMA DOS CUSTOS & A SOLUÇÃO HÍBRIDA
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Viabilidade Econômica: O Dilema de Custos de Amostragem", "GESTÃO DE CUSTOS & LOGÍSTICA DE PESQUISA")

    add_metric_badge(s3, Inches(0.8), Inches(1.45), Inches(3.6), Inches(1.15), "100% Campo (Inviável)", "R$ 75k a 120k", "estimados", C_ROSE)
    add_metric_badge(s3, Inches(4.8), Inches(1.45), Inches(3.6), Inches(1.15), "100% Sintético (Viciado)", "R$ 0,00", "(Rejeição Acadêmica)", C_AMBER)
    add_metric_badge(s3, Inches(8.8), Inches(1.45), Inches(3.7), Inches(1.15), "Estratégia Híbrida SAREL", "Economia ~80%", "(Custo Exequível)", C_EMERALD_LIGHT)

    cards_s3 = [
        ("Cenário A: 100% Campo", [
            ("Demanda Logística:", "Visita presencial a 250-300 pontos dispersos por todo o Paraná."),
            ("Custos Operacionais:", "Combustível, veículos 4x4, diárias de equipe, aluguel de GNSS RTK e 4 a 6 meses de deslocamento contínuo."),
            ("Risco Acadêmico:", "Amostra insuficiente (N < 40 pontos) devido ao orçamento, gerando subajuste (underfitting) no classificador XGBoost."),
        ], C_ROSE),
        ("Cenário B: 100% Sintético", [
            ("Atalho Computacional:", "Gerar dados de erosão e solo a partir de fórmulas matemáticas (RUSLE) ou simulações aleatórias sem checagem empírica."),
            ("Vício Científico:", "O modelo de IA apenas memoriza as equações que o próprio pesquisador inseriu (vazamento de alvo / overfitting trivial)."),
            ("Risco Acadêmico:", "Desqualificação imediata perante a banca examinadora por ausência de dados empíricos observacionais."),
        ], C_AMBER),
        ("Cenário C: Híbrido SAREL", [
            ("Desenho Otimizado:", "Triagem no Earth Engine + Fotointerpretação em Alta Resolução (PlanetScope/VHR) para 80% do dataset."),
            ("Calibração Presencial:", "Subamostra de controle (20%) em campo com KoboToolbox, GNSS RTK e ortomosaicos de drone."),
            ("Rigor Estatístico:", "Matriz de Confusão e Índice Kappa de Cohen (κ ≥ 0,60) medindo e compensando formalmente a taxa de erro humana."),
        ], C_EMERALD_LIGHT),
    ]

    for idx, (title, content, color) in enumerate(cards_s3):
        bx = Inches(0.8 + idx * 4.0)
        add_card(s3, bx, Inches(2.8), Inches(3.7), Inches(3.9), title, content, accent_color=color, title_size=13, body_size=10)

    # =========================================================================
    # SLIDE 4: O PAPEL DO SAREL
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Escopo do Sistema: As Três Funções Centrais do SAREL", "ARQUITETURA FUNCIONAL & ESCOPO DA DISSERTAÇÃO")

    roles = [
        ("1. Propõe Onde Observar", "Desenho Amostral Estratificado & Disperso", [
            "• Executa estratificação multivariada em 18 estratos biofísicos cruzando relevo, solo e exposição.",
            "• Garante formalmente a inclusão de candidatos à classe de controle (Sistema Plantio Direto e pastagem estável).",
            "• Aplica Thinning Geodésico (raio ≥ 1,0 km) para impedir agrupamentos espaciais e pseudorrepetição amostral.",
        ], C_CYAN),
        ("2. Extrai Features Verificáveis", "Proveniência por Variável e Sem Atalhos", [
            "• Coleta séries temporais harmonizadas Sentinel-2 L2A (10 anos) e extrai métricas de exposição de solo nu (Ê).",
            "• Calcula declividade em projeção métrica UTM 22S (SIRGAS 2000) a partir do DEM Copernicus GLO-30.",
            "• Consulta pedologia real via WMS OGC no GeoServer da Embrapa Solos, sem classes inventadas.",
        ], C_EMERALD_LIGHT),
        ("3. Gerencia a Rotulagem Cega", "Campanhas de Validação e Matriz de Treino", [
            "• Exporta cadernos cegos de fotointerpretação, campo (KoboToolbox) e voo de drone, blindando o intérprete.",
            "• Avalia a concordância inter-observador via Kappa de Cohen com alerta bloqueante para κ < 0,60.",
            "• Monta a matriz de treino do XGBoost livre de coordenadas e de variáveis calculadas pelo próprio sistema.",
        ], C_INDIGO),
    ]

    for idx, (head, sub, items, col) in enumerate(roles):
        bx = Inches(0.8 + idx * 4.0)
        bullets = [(sub, "")] + [(item, "") for item in items]
        add_card(s4, bx, Inches(1.55), Inches(3.7), Inches(4.3), head, bullets, accent_color=col, title_size=13.5, body_size=9.8)

    banner = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.7), Inches(0.9))
    banner.fill.solid()
    banner.fill.fore_color.rgb = C_ACCENT_BG
    banner.line.color.rgb = C_EMERALD_LIGHT
    tb_b = s4.shapes.add_textbox(Inches(1.0), Inches(6.05), Inches(11.3), Inches(0.8))
    tf_b = tb_b.text_frame
    tf_b.word_wrap = True
    p_b1 = tf_b.paragraphs[0]
    p_b1.text = "PRINCÍPIO DE DEMARCAÇÃO DO SAREL:"
    p_b1.font.size = Pt(11)
    p_b1.font.bold = True
    p_b1.font.color.rgb = C_EMERALD_DARK
    p_b2 = tf_b.add_paragraph()
    p_b2.text = "O SAREL NÃO classifica erosão e NÃO treina a IA dentro de si. Ele prepara o chão experimental honesto e auditável. O rótulo vem exclusivamente de observadores humanos. O treinamento do XGBoost ocorre externamente."
    p_b2.font.size = Pt(9.5)
    p_b2.font.color.rgb = C_TEXT_DARK

    # =========================================================================
    # SLIDE 5: A LEI FUNDAMENTAL DO SAREL & REGRAS INVIOLÁVEIS
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "A Lei Fundamental do SAREL: Blindagem Pericial Contra Atalhos", "GOVERNANÇA METODOLÓGICA & INTEGRIDADE")

    rules_grid = [
        ("Regra 1 & 5: Dado Real ou Ausência Declarada", [
            ("Proibição de Constantes:", "Terminantemente vedado preencher lacunas com médias ou constantes arbitrárias (ex.: unmask numérico ou defaults em cascata)."),
            ("Tratamento de Nulos:", "O XGBoost lida com ausência de dados nativamente. Ausente é tecnicamente superior a dado fabricado."),
            ("Zero é Valor:", "Declividade 0% ou Ê = 0 são valores físicos legítimos, nunca tratados como ausência."),
        ], C_ROSE),
        ("Regra 3: Proveniência Viaja Junto com o Valor", [
            ("Tríade de Metadados:", "Todo valor numérico é acompanhado de estado, fonte primária oficial e data de aquisição."),
            ("Quatro Selos Formais:", "● medido (sensor primário) • ◊ modelado (equação) • □ tabelado (literatura) • ○ indisponível."),
            ("Rastreabilidade Pericial:", "Possibilita reconstruir a cadeia de evidências de cada linha perante a banca."),
        ], C_EMERALD_LIGHT),
        ("Regra 4: Nada Calculado Vira Rótulo", [
            ("Isolamento de Critérios:", "Scores internos, severidade estimada ou perda RUSLE nunca viram rótulo de treino."),
            ("Prevenção de Vazamento:", "Evita que a IA aprenda limiares artificiais das fórmulas em vez de padrões reais da paisagem."),
            ("Rótulo Exclusivo Humano:", "A classe alvo deriva unicamente de inspeção de campo, fotointerpretação ou drone."),
        ], C_CYAN),
        ("Regra 9: Governança de Decisões (DECISOES.md)", [
            ("Dono do Parâmetro:", "Nenhum limiar, peso ou corte nasce como literal no código sem autor, data e referência bibliográfica."),
            ("Decisão Pendente:", "Enquanto o pesquisador não decidir (ex.: D13, D14, D15), o cálculo gera 'indisponivel' com a causa."),
            ("Imunidade a Arbitrariedades:", "O agente de IA não toma decisões de dissertação no lugar do mestrando."),
        ], C_AMBER),
    ]

    for idx, (title, content, color) in enumerate(rules_grid):
        col = idx % 2
        row = idx // 2
        bx = Inches(0.8 + col * 6.0)
        by = Inches(1.55 + row * 2.65)
        add_card(s5, bx, by, Inches(5.7), Inches(2.45), title, content, accent_color=color, title_size=13, body_size=10)

    # =========================================================================
    # SLIDE 6: DESENHO AMOSTRAL & THINNING GEODÉSICO
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Desenho Amostral: 18 Estratos Multivariados & Thinning Geodésico", "DESENHO EXPERIMENTAL GEOESPACIAL")

    sample_steps = [
        ("1. Estratificação Multivariada", "18 Estratos (Ŝ x Ê x K̂)", [
            ("Dimensão Ŝ (Relevo):", "Tercis de declividade (Suave, Moderada, Forte) derivados do DEM Copernicus."),
            ("Dimensão Ê (Solo Nu):", "Tercis de frequência multitemporal de solo descoberto da série Sentinel-2."),
            ("Dimensão K̂ (Pedologia):", "2 níveis de erodibilidade derivados do levantamento pedológico da Embrapa."),
        ], C_CYAN),
        ("2. Thinning Geodésico", "Raio Mínimo d ≥ 1,0 km (P02)", [
            ("Primeira Lei de Tobler:", "Pixels contíguos possuem solos e relevo similares, causando inflação artificial de acurácia."),
            ("Algoritmo Fisher-Yates:", "Espaçamento euclidiano/haversine determinístico com semente registrada (P07)."),
            ("Dispersão Homogênea:", "Força a representatividade regional em toda a extensão da bacia hidrográfica."),
        ], C_EMERALD_LIGHT),
        ("3. Garantia de Controles", "Balanceamento de Classes (0 e 1)", [
            ("Classe 1 (Erosão):", "Candidatos estratificados com histórico de exposição e gradiente de rampa."),
            ("Classe 0 (Controle):", "Amostragem em áreas sob Sistema Plantio Direto (SPD) e pastagens estáveis."),
            ("Prevenção do 'Mock':", "Garante que o modelo aprenda o contraste biofísico real entre equilíbrio e degradação."),
        ], C_AMBER),
    ]

    for idx, (title, sub, bullets, color) in enumerate(sample_steps):
        bx = Inches(0.8 + idx * 4.0)
        content = [(sub, "")] + bullets
        add_card(s6, bx, Inches(1.55), Inches(3.7), Inches(5.15), title, content, accent_color=color, title_size=13.5, body_size=10.2)

    # =========================================================================
    # SLIDE 7: SENSORIAMENTO REMOTO & HARMONIZAÇÃO TEMPORAL
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Sensoriamento Remoto: Sentinel-2, DEM Copernicus & Decomposição Harmônica", "EXTRAÇÃO DE PREDITORES FÍSICO-INFORMADOS")

    remote_cards = [
        ("Sentinel-2 MSI L2A (BOA)", [
            ("Bandas Físicas:", "B2 (Azul), B4 (Vermelho), B8 (NIR) e B11/B12 (SWIR) com reflectância de fundo de atmosfera."),
            ("Índices Espectrais:", "NDVI (vigor vegetal) e BSI (Bare Soil Index) computados pixel a pixel na série temporal."),
            ("Máscara de Nuvens SCL:", "Filtro Scene Classification Layer com descarte de pixels nublados sem reposição de constantes."),
        ], C_CYAN),
        ("Copernicus DEM GLO-30", [
            ("Projeção Métrica Oficial:", "Cálculo de declividade e elevação em SIRGAS 2000 / UTM 22S (EPSG:31982)."),
            ("Correção do Erro 3857:", "Elimina a subestimação sistemática de 10% na declividade causada pelo Web Mercator no Paraná."),
            ("Atributos de Terreno:", "Elevação ortométrica, declividade (graus/%), curvatura de vertente e acúmulo de fluxo."),
        ], C_EMERALD_LIGHT),
        ("Decomposição Harmônica OLS", [
            ("Ajuste Multivariado:", "Regressão harmônica de 1 e 2 ciclos anuais (Zhu & Woodcock 2014) para cada banda espectral."),
            ("Parâmetros Extraídos:", "Tendência temporal (b1), amplitudes anual/semianual e fases espectrais de cada comprimento de onda."),
            ("Métricas de Qualidade:", "Cada coeficiente armazena número de observações válidas, R² do ajuste e erro-padrão."),
        ], C_INDIGO),
        ("Dinâmica do Solo Exposto (Ê)", [
            ("Frequência de Solo Nu (Ê):", "Fração temporal em que o solo permaneceu descoberto nos anos de monitoramento."),
            ("Sequência Contínua:", "Maior número consecutivo de meses sob exposição crítica aos agentes erosivos."),
            ("Mês Modal de Exposição:", "Identifica o período do ano com maior vulnerabilidade climática no Paraná."),
        ], C_AMBER),
    ]

    for idx, (title, bullets, color) in enumerate(remote_cards):
        col = idx % 2
        row = idx // 2
        bx = Inches(0.8 + col * 6.0)
        by = Inches(1.55 + row * 2.65)
        add_card(s7, bx, by, Inches(5.7), Inches(2.45), title, bullets, accent_color=color, title_size=13, body_size=10)

    # =========================================================================
    # SLIDE 8: LINHA DE BASE RUSLE REGIONAL
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Linha de Base RUSLE: Instrumento Comparativo com Respaldo Científico", "MODELAGEM FÍSICO-MATEMÁTICA CONVENCIONAL")

    f_box = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.7), Inches(0.95))
    f_box.fill.solid()
    f_box.fill.fore_color.rgb = C_ACCENT_BG
    f_box.line.color.rgb = C_EMERALD_LIGHT
    tb_f = s8.shapes.add_textbox(Inches(1.0), Inches(1.55), Inches(11.3), Inches(0.85))
    tf_f = tb_f.text_frame
    tf_f.word_wrap = True
    p_f = tf_f.paragraphs[0]
    p_f.text = "A = R × K × LS × C × P   [Estimativa de Perda de Solo em t / (ha · ano)]"
    p_f.font.size = Pt(17)
    p_f.font.bold = True
    p_f.font.color.rgb = C_EMERALD_DARK
    p_f.alignment = PP_ALIGN.CENTER

    rusle_cards = [
        ("Fator C: Durigon et al. (2014) [DECIDIDO - D01]", [
            ("Equação Regional:", "C = (1 − NDVI) / 2 calibrada para bacias tropicais brasileiras sem parâmetros livres."),
            ("Sem Cortes Disfarçados:", "NDVI estritamente contido em [-1, 1], garantindo C contido em [0, 1] sem Math.min/max artificiais."),
            ("Análise de Sensibilidade:", "Comparação programada com a formulação de van der Knijff et al. (2000)."),
        ], C_EMERALD_LIGHT),
        ("Fator P: Práticas Conservacionistas [TABELADO]", [
            ("Valor Referencial:", "Fator P = 1,0 quando as práticas mecânicas do imóvel forem desconhecidas (Renard et al., 1997)."),
            ("Proveniência Declarada:", "Registrado estritamente como 'tabelado', nunca apresentado à banca como se fosse medição de campo."),
            ("Terraceamento:", "Ajustável na rotulagem presencial caso terraços de retenção íntegros sejam comprovados."),
        ], C_CYAN),
        ("Fatores R, K e LS [PENDENTES - D13, D14, D15]", [
            ("Trava Inviolável (Regra 9):", "Enquanto as equações não forem formalmente selecionadas pelo pesquisador, permanecem pendentes."),
            ("Bloqueio de Perda de Solo:", "Perda de solo fica 'indisponivel (decisao-pendente)'. Proibido usar multiplicadores inventados."),
            ("Invariante 1 Ativo:", "A perda só é exportada se os 5 fatores e a memória de cálculo com a equação estiverem preenchidos."),
        ], C_AMBER),
    ]

    for idx, (title, bullets, color) in enumerate(rusle_cards):
        bx = Inches(0.8 + idx * 4.0)
        add_card(s8, bx, Inches(2.65), Inches(3.7), Inches(4.05), title, bullets, accent_color=color, title_size=12.5, body_size=9.8)

    # =========================================================================
    # SLIDE 9: INTELIGÊNCIA FUNDIÁRIA & CONFORMIDADE LGPD
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Inteligência Fundiária: Integração SICAR/SIGEF/SNCR & Proteção LGPD", "BASE TERRITORIAL AUDITÁVEL & LGPD")

    bases_s9 = [
        ("SICAR / MMA", "559.899 Imóveis Rurais no PR", [
            ("Código CAR Oficial:", "Identificador único federal de cada imóvel cadastrado no Paraná."),
            ("Atributos Cartográficos:", "Área total em hectares, módulos fiscais e situação cadastral do imóvel rural."),
            ("Perímetro Vetorial:", "Polígono georreferenciado da propriedade para conferência de limites e confrontantes."),
        ], C_EMERALD_LIGHT),
        ("SIGEF / INCRA", "170.046 Parcelas Certificadas", [
            ("Lei 10.267:", "Certificação de georreferenciamento de imóveis rurais com alta precisão métrica."),
            ("Segurança Jurídica:", "Identificação do número de Matrícula no Cartório de Registro de Imóveis (CRI)."),
            ("Anotação Técnica:", "Registro de ART/CREA do engenheiro agrimensor responsável pela demarcação."),
        ], C_CYAN),
        ("SNCR / Receita Federal", "957.183 Cadastros de Titulares", [
            ("Viabilização do Acesso:", "Permite que a equipe de campo contate o produtor para autorização de entrada na fazenda."),
            ("Cruzamento Local R*Tree:", "Consultas espaciais executadas em menos de 5ms no banco SQLite local."),
            ("Total Imunidade:", "Zero divergências cadastrais atestadas na auditoria contra os dados abertos oficiais."),
        ], C_INDIGO),
    ]

    for idx, (title, vol, bullets, color) in enumerate(bases_s9):
        bx = Inches(0.8 + idx * 4.0)
        content = [(vol, "")] + bullets
        add_card(s9, bx, Inches(1.55), Inches(3.7), Inches(3.9), title, content, accent_color=color, title_size=14, body_size=10)

    lgpd_card = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.65), Inches(11.7), Inches(1.25))
    lgpd_card.fill.solid()
    lgpd_card.fill.fore_color.rgb = C_ACCENT_BG
    lgpd_card.line.color.rgb = C_EMERALD_DARK
    tb_l = s9.shapes.add_textbox(Inches(1.0), Inches(5.7), Inches(11.3), Inches(1.1))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    p_l1 = tf_l.paragraphs[0]
    p_l1.text = "CONFORMIDADE RIGOROSA COM A LGPD (LEI 13.709/2018, ART. 7º, IV):"
    p_l1.font.size = Pt(11)
    p_l1.font.bold = True
    p_l1.font.color.rgb = C_EMERALD_DARK
    p_l2 = tf_l.add_paragraph()
    p_l2.text = "Os dados fundiários servem exclusivamente à gestão operacional de acesso em campo e nunca entram na matriz de treino do XGBoost. Nomes e documentos de titulares são mascarados byte a byte pela máscara oficial do SNCR. Rotas fundiárias e tokens transitam em sessões efêmeras no servidor, com recusa automática a conexões externas."
    p_l2.font.size = Pt(9.5)
    p_l2.font.color.rgb = C_TEXT_DARK

    # =========================================================================
    # SLIDE 10: PROTOCOLO DE ROTULAGEM CEGA & KAPPA DE COHEN
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Protocolo de Rotulagem Cega: Controle de Viés & Validação Estatística", "CAMPANHAS DE ROTULAGEM & CONCORDÂNCIA HUMANA")

    phases_cards = [
        ("Fase A: Interpretação Visual Cega", [
            ("Imagens PlanetScope/VHR:", "Inspeção em resolução de 3m/30cm com data de aquisição visível no visor."),
            ("Isolamento Cego:", "O intérprete desconhece predições do modelo, fatores RUSLE, declividade e estrato."),
            ("Critério Padronizado:", "Avaliação de cobertura vegetal, decapitação de solo, carreadores e terraços."),
        ], C_CYAN),
        ("Fase B: Validação em Campo (Kobo)", [
            ("Subamostra Presencial:", "Equipe visita ~20% dos pontos estratificados munida de GNSS RTK e KoboToolbox."),
            ("Cegueira Cruzada:", "A equipe de campo desconhece o rótulo atribuído na Fase A, medindo a taxa de erro humana."),
            ("Marcadores Biofísicos:", "Avaliação milimétrica de pedestais de erosão, descalçamento radicular e crosta."),
        ], C_EMERALD_LIGHT),
        ("Fase C: Concordância Kappa", [
            ("Kappa de Cohen (1960):", "Cálculo rigoroso de concordância inter-observador ajustado ao acaso."),
            ("Alerta Bloqueante (κ < 0,60):", "Se a concordância for inferior a substancial (Landis & Koch), a matriz é travada."),
            ("Resolução de Divergência:", "Pontos divergentes são arbitrados por perito sênior ou descartados da matriz."),
        ], C_AMBER),
        ("Fase D: Drone Held-Out", [
            ("Ortomosaicos de VANT:", "Voo de drone gerando modelos tridimensionais centimétricos nos sítios padrão-ouro."),
            ("Segregação Obrigatória:", "As amostras de drone são mantidas compulsoriamente como conjunto de teste (held-out)."),
            ("Testagem Independente:", "Nunca entram no treino do XGBoost, garantindo avaliação final não contaminada."),
        ], C_INDIGO),
    ]

    for idx, (title, bullets, color) in enumerate(phases_cards):
        bx = Inches(0.8 + idx * 3.0)
        add_card(s10, bx, Inches(1.55), Inches(2.75), Inches(5.15), title, bullets, accent_color=color, title_size=13, body_size=9.8)

    # =========================================================================
    # SLIDE 11: MODELAGEM PREDITIVA XGBoost LOCO & SHAP
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Modelagem Preditiva: XGBoost com Validação Cruzada Espacial LOCO", "APRENDIZADO DE MÁQUINA & EXPLICABILIDADE FÍSICA")

    xgb_cards = [
        ("Por que LOCO e não K-Fold Tradicional?", [
            ("Autocorrelação Espacial:", "Amostras geoespaciais próximas violam a premissa i.i.d. O K-Fold comum mistura pixels vizinhos, inflacionando a acurácia."),
            ("Leave-One-Catchment-Out:", "O algoritmo treina em 4 macrobacias hidrográficas do Paraná e é testado na 5ª bacia inteiramente omitida."),
            ("Capacidade de Generalização:", "Mede a eficácia real do método em extrapolar a predição para novos territórios sem vazamento de vizinhança."),
        ], C_ROSE),
        ("Explicabilidade Física via SHAP", [
            ("Valores de Shapley:", "Mede a contribuição marginal de cada variável preditora (NDVI, BSI, Declividade, K, R) na decisão do modelo."),
            ("Auditoria de Coerência:", "Garante que a erosão seja atribuída a processos físicos reais (alta declividade + solo exposto), e não a ruídos estatísticos."),
            ("Entregáveis da Dissertação:", "Gráficos SHAP Summary (Beeswarm Plot) e gráficos de dependência parcial revelando limiares críticos."),
        ], C_CYAN),
        ("Métricas Científicas Reais Esperadas", [
            ("Acurácia Realista (84% a 93%):", "Ambientes naturais e sensores orbitais possuem ruídos. Resultados de 100% decorrem de vícios e foram barrados."),
            ("ROC-AUC Esperada (0,88 a 0,95):", "Desempenho de alta separabilidade com calibração probabilística rigorosa."),
            ("Integridade Assegurada:", "O script Python treinar_xgboost_loco.py possui trava anti-mock ativa que exige dados reais."),
        ], C_EMERALD_LIGHT),
    ]

    for idx, (title, bullets, color) in enumerate(xgb_cards):
        bx = Inches(0.8 + idx * 4.0)
        add_card(s11, bx, Inches(1.55), Inches(3.7), Inches(5.15), title, bullets, accent_color=color, title_size=13, body_size=10.2)

    # =========================================================================
    # SLIDE 12: INVARIANTES DE EXPORTAÇÃO & BLINDAGEM PERICIAL
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Invariantes de Exportação: 7 Portões de Segurança Contra Falhas", "CONTROLE DE QUALIDADE AUTOMATIZADO")

    invariants_list = [
        ("Invariante 1", "Coerência RUSLE", "Perda de solo preenchida ⟺ 5 fatores preenchidos ⟺ memória de cálculo válida presente.", C_EMERALD_LIGHT),
        ("Invariante 2", "Lista de Permissão do Perfil", "Colunas contidas estritamente no perfil cego (pesquisador, campo, voo ou matriz de treino).", C_CYAN),
        ("Invariante 3", "Rastreamento de Campos Estimados", "A coluna Campos_Estimados lista exatamente as variáveis com proveniência diferente de 'medido'.", C_INDIGO),
        ("Invariante 4", "Rastreabilidade Orbital", "Dado derivado de satélite exige PRODUCT_ID das cenas, versão do motor e data de processamento.", C_AMBER),
        ("Invariante 5", "Afirmação Negativa Fundiária", "Declaração 'sem-correspondencia' só é aceita após consulta cadastral bem-sucedida sem match.", C_ROSE),
        ("Invariante 6", "Lista Negra de Literais Geográficos", "Bloqueia termos de desenvolvimento ('Custom', 'Bacia Local') ou nome de UF na coluna município.", C_EMERALD_DARK),
        ("Invariante 7", "Detector de Constante Disfarçada", "Nenhuma coluna numérica tem valor idêntico em todas as linhas (N ≥ 21), impedindo vícios legados.", C_ROSE),
    ]

    for idx, (inv_id, inv_nome, inv_desc, col) in enumerate(invariants_list):
        by = Inches(1.55 + idx * 0.76)
        box = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), by, Inches(11.7), Inches(0.70))
        box.fill.solid()
        box.fill.fore_color.rgb = C_WHITE
        box.line.color.rgb = C_BORDER_LIGHT
        box.line.width = Pt(1)

        tag = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), by + Inches(0.12), Inches(1.8), Inches(0.46))
        tag.fill.solid()
        tag.fill.fore_color.rgb = col
        tag.line.color.rgb = col
        p_t = tag.text_frame.paragraphs[0]
        p_t.text = f"{inv_id.upper()}"
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = C_WHITE
        p_t.alignment = PP_ALIGN.CENTER

        tb = s12.shapes.add_textbox(Inches(2.85), by + Inches(0.08), Inches(9.5), Inches(0.55))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        r1 = p1.add_run()
        r1.text = inv_nome + ": "
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = C_TEXT_DARK
        r2 = p1.add_run()
        r2.text = inv_desc
        r2.font.size = Pt(10)
        r2.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 13: CONCLUSÃO & CONTRIBUIÇÃO CIENTÍFICA
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13, C_NAVY_DARK)

    tb_c13 = s13.shapes.add_textbox(Inches(1.0), Inches(1.1), Inches(11.333), Inches(5.6))
    tf_c13 = tb_c13.text_frame
    tf_c13.word_wrap = True

    p0 = tf_c13.paragraphs[0]
    p0.text = "CONCLUSÃO & CONTRIBUIÇÕES PARA O AGRONEGÓCIO PARANAENSE"
    p0.font.size = Pt(12)
    p0.font.bold = True
    p0.font.color.rgb = C_EMERALD_LIGHT

    p1 = tf_c13.add_paragraph()
    p1.text = "Um Instrumento Computacional Ético, Auditável e Replicável"
    p1.font.size = Pt(25)
    p1.font.bold = True
    p1.font.color.rgb = C_WHITE
    p1.space_before = Pt(8)

    concl_points = [
        ("Viabilidade Prática Comprovada:", "A estratégia híbrida reduz os custos operacionais de amostragem em ~80%, tornando viável a validação de modelos de IA para conservação de solos no mestrado."),
        ("Ruptura com o Dado Presumido:", "O SAREL encerra a era dos valores fabricados por atalhos computacionais. O sistema prefere declarar a ausência a exibir uma medição ilusória à banca."),
        ("Rigor Metodológico Internacional:", "Estratificação multivariada, Thinning Geodésico, validação cruzada espacial LOCO e explicabilidade SHAP alinham a pesquisa às melhores práticas da literatura de sensoriamento remoto."),
        ("Prontidão para a Campanha Real:", "A infraestrutura computacional (Next.js, TypeScript, Python, SQLite e GEE) está 100% testada e homologada, pronta para a execução da campanha de campo e fotointerpretação."),
    ]

    for title, desc in concl_points:
        p = tf_c13.add_paragraph()
        r1 = p.add_run()
        r1.text = f"• {title} "
        r1.font.bold = True
        r1.font.size = Pt(11.5)
        r1.font.color.rgb = C_EMERALD_LIGHT
        r2 = p.add_run()
        r2.text = desc
        r2.font.size = Pt(11)
        r2.font.color.rgb = C_TEXT_LIGHT
        p.space_before = Pt(10)

    p_end = tf_c13.add_paragraph()
    p_end.text = "PPGTCA 2026 • Pesquisador: Luis Alfredo • Programa de Pós-Graduação em Tecnologias Computacionais para o Agronegócio"
    p_end.font.size = Pt(10.5)
    p_end.font.bold = True
    p_end.font.color.rgb = C_CYAN
    p_end.space_before = Pt(24)

    # Save
    prs.save(output_pptx_path)
    print(f"\n[SUCESSO] Apresentação SAREL v2 gerada com sucesso!")
    print(f"Arquivo gerado: {os.path.abspath(output_pptx_path)}")
    print(f"Total de slides criados: {len(prs.slides)}\n")

if __name__ == "__main__":
    out = "Apresentacao_SAREL_PPGTCA_2026.pptx"
    if len(sys.argv) > 1:
        out = sys.argv[1]
    create_deck(out)
